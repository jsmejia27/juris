# case_vault_service.py
import os
import io
import json
import time
import uuid
import logging
from typing import List, Dict, Any, Optional
from datetime import datetime

from qdrant_client import QdrantClient
from qdrant_client.http import models
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_ollama import OllamaEmbeddings, OllamaLLM

from rag_pipeline import (
    get_shared_qdrant_client,
    _QDRANT_LOCK,
    DEFAULT_OLLAMA_URL,
    DEFAULT_STORAGE_DIR,
    LegalRAGPipeline
)

logger = logging.getLogger(__name__)

VAULT_BASE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "vault", "cases")
VAULT_COLLECTION = "juris_case_vault"

# Supported text extractors
def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract clean text from PDF bytes using pypdf."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                pages_text.append(f"--- [Page {i+1}] ---\n{text.strip()}")
        return "\n\n".join(pages_text)
    except Exception as e:
        logger.error(f"Failed to extract PDF text: {e}")
        return ""

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract clean text from DOCX bytes using python-docx."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    paragraphs.append(row_text)
                    
        return "\n\n".join(paragraphs)
    except Exception as e:
        logger.error(f"Failed to extract DOCX text: {e}")
        return ""

def extract_document_text(filename: str, file_bytes: bytes) -> str:
    """Extract text based on file extension."""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_bytes)
    else:
        # Fallback to UTF-8 / ASCII text
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1", errors="replace")


class CaseVaultService:
    def __init__(
        self,
        ollama_url: str = DEFAULT_OLLAMA_URL,
        embed_model: str = "nomic-embed-text"
    ):
        self.ollama_url = ollama_url
        self.embed_model = embed_model
        self.client = get_shared_qdrant_client()
        os.makedirs(VAULT_BASE_DIR, exist_ok=True)
        self._init_vault_collection()

        self.embedder = OllamaEmbeddings(
            base_url=self.ollama_url,
            model=self.embed_model
        )
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=150,
            separators=["\n\n", "\n", ". ", "; ", " ", ""]
        )

    def _init_vault_collection(self):
        """Ensure Qdrant collection for Case Vault exists."""
        with _QDRANT_LOCK:
            try:
                collections = self.client.get_collections().collections
                exists = any(c.name == VAULT_COLLECTION for c in collections)
                if not exists:
                    self.client.create_collection(
                        collection_name=VAULT_COLLECTION,
                        vectors_config=models.VectorParams(
                            size=768,
                            distance=models.Distance.COSINE
                        )
                    )
                    # Index case_id payload field for fast filtering
                    self.client.create_payload_index(
                        collection_name=VAULT_COLLECTION,
                        field_name="case_id",
                        field_schema=models.PayloadSchemaType.KEYWORD
                    )
                    logger.info(f"Initialized Qdrant collection: {VAULT_COLLECTION}")
            except Exception as e:
                logger.error(f"Error checking/creating vault collection {VAULT_COLLECTION}: {e}")

    # =========================================================================
    # DOSSIER CRUD MANAGEMENT
    # =========================================================================

    def list_cases(self) -> List[Dict[str, Any]]:
        """List all active case dossiers with summary metadata."""
        cases = []
        if not os.path.exists(VAULT_BASE_DIR):
            return cases

        for case_id in os.listdir(VAULT_BASE_DIR):
            case_path = os.path.join(VAULT_BASE_DIR, case_id)
            meta_file = os.path.join(case_path, "meta.json")
            if os.path.isdir(case_path) and os.path.exists(meta_file):
                try:
                    with open(meta_file, "r", encoding="utf-8") as f:
                        meta = json.load(f)
                        cases.append(meta)
                except Exception as e:
                    logger.warning(f"Failed to read meta for case {case_id}: {e}")

        # Sort by updated_at descending
        cases.sort(key=lambda x: x.get("updated_at", 0), reverse=True)
        return cases

    def get_case(self, case_id: str) -> Optional[Dict[str, Any]]:
        """Retrieve full details and document records for a specific case dossier."""
        case_path = os.path.join(VAULT_BASE_DIR, case_id)
        meta_file = os.path.join(case_path, "meta.json")
        if os.path.exists(meta_file):
            with open(meta_file, "r", encoding="utf-8") as f:
                return json.load(f)
        return None

    def create_case(
        self,
        title: str,
        docket_number: str = "",
        court: str = "",
        case_type: str = "Civil Litigation",
        parties: str = "",
        description: str = ""
    ) -> Dict[str, Any]:
        """Create a new local case dossier workspace."""
        case_id = f"CASE-{uuid.uuid4().hex[:8].upper()}"
        case_dir = os.path.join(VAULT_BASE_DIR, case_id)
        docs_dir = os.path.join(case_dir, "documents")
        os.makedirs(docs_dir, exist_ok=True)

        now = time.time()
        meta = {
            "case_id": case_id,
            "title": title.strip(),
            "docket_number": docket_number.strip(),
            "court": court.strip(),
            "case_type": case_type.strip(),
            "parties": parties.strip(),
            "description": description.strip(),
            "created_at": now,
            "updated_at": now,
            "created_date_str": datetime.fromtimestamp(now).strftime("%Y-%m-%d %H:%M:%S"),
            "document_count": 0,
            "total_chunks": 0,
            "documents": []
        }

        with open(os.path.join(case_dir, "meta.json"), "w", encoding="utf-8") as f:
            json.dump(meta, f, indent=2)

        logger.info(f"Created new Case Dossier: {case_id} - {title}")
        return meta

    def delete_case(self, case_id: str) -> bool:
        """Delete case files and purge indexed vectors from Qdrant."""
        case_dir = os.path.join(VAULT_BASE_DIR, case_id)
        
        # 1. Delete Qdrant vectors for this case
        with _QDRANT_LOCK:
            try:
                self.client.delete(
                    collection_name=VAULT_COLLECTION,
                    points_selector=models.FilterSelector(
                        filter=models.Filter(
                            must=[
                                models.FieldCondition(
                                    key="case_id",
                                    match=models.MatchValue(value=case_id)
                                )
                            ]
                        )
                    )
                )
                logger.info(f"Purged vectors from Qdrant for case {case_id}")
            except Exception as e:
                logger.error(f"Error purging Qdrant vectors for {case_id}: {e}")

        # 2. Delete local files
        if os.path.exists(case_dir):
            import shutil
            shutil.rmtree(case_dir, ignore_errors=True)
            logger.info(f"Deleted local directory for case {case_id}")
            return True
        return False

    # =========================================================================
    # DOCUMENT INGESTION & INDEXING
    # =========================================================================

    def add_document_to_case(
        self,
        case_id: str,
        filename: str,
        file_bytes: bytes,
        doc_category: str = "Pleading / Motion"
    ) -> Dict[str, Any]:
        """Parse, save, and index a document inside a case dossier."""
        case_meta = self.get_case(case_id)
        if not case_meta:
            raise ValueError(f"Case dossier {case_id} not found.")

        case_dir = os.path.join(VAULT_BASE_DIR, case_id)
        docs_dir = os.path.join(case_dir, "documents")
        os.makedirs(docs_dir, exist_ok=True)

        doc_id = f"DOC-{uuid.uuid4().hex[:6].upper()}"
        safe_filename = f"{doc_id}_{filename}"
        file_path = os.path.join(docs_dir, safe_filename)

        # 1. Save original file
        with open(file_path, "wb") as f:
            f.write(file_bytes)

        # 2. Extract Text
        extracted_text = extract_document_text(filename, file_bytes)
        if not extracted_text.strip():
            extracted_text = "[No text could be extracted from this document.]"

        # 3. Chunk text
        chunks = self.splitter.split_text(extracted_text)
        if not chunks:
            chunks = [extracted_text]

        # 4. Generate Embeddings and Index to Qdrant
        points = []
        now = time.time()
        for idx, chunk in enumerate(chunks):
            point_id = str(uuid.uuid4())
            try:
                vector = self.embedder.embed_query(chunk)
            except Exception as e:
                logger.warning(f"Embedding error: {e}. Utilizing zero vector fallback.")
                vector = [0.0] * 768

            points.append(
                models.PointStruct(
                    id=point_id,
                    vector=vector,
                    payload={
                        "case_id": case_id,
                        "doc_id": doc_id,
                        "doc_name": filename,
                        "doc_category": doc_category,
                        "chunk_index": idx,
                        "total_chunks": len(chunks),
                        "text": chunk,
                        "timestamp": now
                    }
                )
            )

        with _QDRANT_LOCK:
            try:
                self.client.upsert(
                    collection_name=VAULT_COLLECTION,
                    points=points
                )
            except Exception as e:
                logger.error(f"Failed to upsert points to Qdrant for {doc_id}: {e}")

        # 5. Update case metadata
        doc_entry = {
            "doc_id": doc_id,
            "filename": filename,
            "saved_filename": safe_filename,
            "category": doc_category,
            "file_size": len(file_bytes),
            "char_count": len(extracted_text),
            "chunk_count": len(chunks),
            "uploaded_at": now,
            "date_str": datetime.fromtimestamp(now).strftime("%Y-%m-%d %H:%M:%S")
        }

        case_meta["documents"].append(doc_entry)
        case_meta["document_count"] = len(case_meta["documents"])
        case_meta["total_chunks"] += len(chunks)
        case_meta["updated_at"] = now

        with open(os.path.join(case_dir, "meta.json"), "w", encoding="utf-8") as f:
            json.dump(case_meta, f, indent=2)

        return doc_entry

    # =========================================================================
    # CASE RAG & CROSS-EXAMINATION ENGINE
    # =========================================================================

    def search_case_vault(
        self,
        case_id: str,
        query: str,
        top_k: int = 5
    ) -> List[Dict[str, Any]]:
        """Retrieve the most relevant factual chunks from the case dossier."""
        try:
            query_vector = self.embedder.embed_query(query)
        except Exception:
            query_vector = [0.0] * 768

        with _QDRANT_LOCK:
            try:
                response = self.client.query_points(
                    collection_name=VAULT_COLLECTION,
                    query=query_vector,
                    query_filter=models.Filter(
                        must=[
                            models.FieldCondition(
                                key="case_id",
                                match=models.MatchValue(value=case_id)
                            )
                        ]
                    ),
                    limit=top_k
                )
                search_results = response.points
            except Exception as e:
                logger.error(f"Search error on {VAULT_COLLECTION}: {e}")
                search_results = []

        results = []
        for hit in search_results:
            results.append({
                "doc_name": hit.payload.get("doc_name", "Unknown"),
                "doc_category": hit.payload.get("doc_category", "Pleading"),
                "chunk_index": hit.payload.get("chunk_index", 0),
                "score": round(float(hit.score), 4),
                "text": hit.payload.get("text", "")
            })
        return results

    def cross_examine_case(
        self,
        case_id: str,
        focus_issue: str,
        llm_model: str = "qwen3.5:9b"
    ) -> Dict[str, Any]:
        """
        Cross-examine case facts against Philippine jurisprudence:
        1. Retrieves factual evidence from the Case Vault.
        2. Retrieves controlling Supreme Court doctrines & statutes from philippine_law.
        3. Synthesizes an adversarial cross-examination report highlighting discrepancies,
           procedural defects, witness credibility risks, and controlling SC doctrines.
        """
        case_meta = self.get_case(case_id)
        if not case_meta:
            raise ValueError(f"Case {case_id} not found.")

        # 1. Retrieve Case Facts
        case_chunks = self.search_case_vault(case_id, focus_issue, top_k=6)
        case_facts_text = "\n\n".join([
            f"[Source: {c['doc_name']} ({c['doc_category']})]\n{c['text']}"
            for c in case_chunks
        ])

        # 2. Retrieve Philippine Legal Authorities
        legal_pipeline = LegalRAGPipeline(llm_model=llm_model)
        legal_docs = legal_pipeline.retrieve_legal_context(focus_issue, limit=4)
        legal_context_text = "\n\n".join([
            f"[Authority: {d.get('title') if isinstance(d, dict) else getattr(d, 'metadata', {}).get('title', 'Supreme Court Decision')} ({d.get('year') if isinstance(d, dict) else getattr(d, 'metadata', {}).get('year', 'N/A')}) - GR: {(d.get('gr_no') or d.get('gr_number')) if isinstance(d, dict) else getattr(d, 'metadata', {}).get('gr_number', 'N/A')}]\n{d.get('text', '') if isinstance(d, dict) else getattr(d, 'page_content', '')}"
            for d in legal_docs
        ])

        # 3. Formulate Cross-Examination Prompt
        prompt = f"""You are **Juris Senior Litigation Counsel**, an elite Philippine trial attorney and legal strategist.

Analyze the provided **CASE RECORD FACTS & WITNESS STATEMENTS** against the **CONTROLLING PHILIPPINE JURISPRUDENCE & LAWS**.

===============================================================================
CASE DOSSIER: {case_meta.get('title')} (Docket No: {case_meta.get('docket_number', 'N/A')})
COURT: {case_meta.get('court', 'Regional Trial Court / NLRC')} | PARTIES: {case_meta.get('parties', 'N/A')}
FOCUS ISSUE: {focus_issue}
===============================================================================

--------------------------------------------------
SECTION A: RECORDED CASE EVIDENCE & STATEMENTS
--------------------------------------------------
{case_facts_text if case_facts_text else "[No documents found matching focus issue in vault.]"}

--------------------------------------------------
SECTION B: GOVERNING PHILIPPINE JURISPRUDENCE & LAWS
--------------------------------------------------
{legal_context_text if legal_context_text else "[No controlling statutes retrieved.]"}

===============================================================================
INSTRUCTIONS FOR LITIGATION STRATEGY & CROSS-EXAMINATION MEMORANDUM
===============================================================================
Generate a structured, rigorous **Cross-Examination & Case Assessment Memorandum** in Markdown with these exact sections:

### I. EXECUTIVE LITIGATION SUMMARY
Provide a crisp strategic assessment of the legal strengths and vulnerabilities of this case under Philippine law.

### II. FACTUAL CONTRADICTIONS & WITNESS CREDIBILITY AUDIT
Identify any discrepancies, self-serving claims, temporal impossibilities, or missing evidentiary links in the record.

### III. CONTROLLING SUPREME COURT DOCTRINES APPLIED
Analyze the case facts under the governing Supreme Court precedents retrieved in Section B (cite specific G.R. numbers, ponentes, and doctrine currency).

### IV. TARGETED CROSS-EXAMINATION QUESTIONS FOR OPPOSING WITNESSES
Provide 5 to 7 sharp, leading questions formatted for trial or deposition cross-examination to lock the opposing witness into concessions or expose contradictions.

### V. PROCEDURAL & SUBSTANTIVE ACTION PLAN
Specific motions, objections, or documentary exhibits to introduce to secure victory.
"""

        try:
            llm = OllamaLLM(
                base_url=self.ollama_url,
                model=llm_model,
                temperature=0.15
            )
            analysis_text = llm.invoke(prompt)
        except Exception as e:
            logger.error(f"LLM generation failed: {e}")
            analysis_text = f"Cross-examination analysis could not be completed: {e}"

        return {
            "case_id": case_id,
            "case_title": case_meta.get("title"),
            "focus_issue": focus_issue,
            "analysis": analysis_text,
            "case_sources": case_chunks,
            "legal_authorities": [
                {
                    "title": d.get("title") if isinstance(d, dict) else getattr(d, "metadata", {}).get("title", ""),
                    "gr_number": (d.get("gr_no") or d.get("gr_number", "")) if isinstance(d, dict) else getattr(d, "metadata", {}).get("gr_number", ""),
                    "year": d.get("year") if isinstance(d, dict) else getattr(d, "metadata", {}).get("year", ""),
                    "status": (d.get("doctrine_status") or "Active Precedent") if isinstance(d, dict) else getattr(d, "metadata", {}).get("doctrine_status", "Active Precedent")
                }
                for d in legal_docs
            ]
        }

# tests/test_case_vault.py
import pytest
import os
import sys
import shutil

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from case_vault_service import (
    CaseVaultService,
    extract_text_from_pdf,
    extract_text_from_docx,
    extract_document_text,
    VAULT_BASE_DIR
)

@pytest.fixture
def temp_vault_service(tmp_path, monkeypatch):
    test_vault_dir = str(tmp_path / "cases")
    monkeypatch.setattr("case_vault_service.VAULT_BASE_DIR", test_vault_dir)
    service = CaseVaultService()
    return service

def test_case_dossier_lifecycle(temp_vault_service):
    # 1. Create Case
    case = temp_vault_service.create_case(
        title="People v. Juan Dela Cruz",
        docket_number="Crim. Case No. 2026-10492",
        court="RTC Branch 14, Manila",
        case_type="Criminal",
        parties="People of the Philippines vs. Juan Dela Cruz",
        description="Alleged violation of Cybercrime Prevention Act (RA 10175)"
    )
    assert case["case_id"].startswith("CASE-")
    assert case["title"] == "People v. Juan Dela Cruz"
    assert case["document_count"] == 0

    # 2. List Cases
    cases = temp_vault_service.list_cases()
    assert len(cases) == 1
    assert cases[0]["case_id"] == case["case_id"]

    # 3. Get Case
    fetched = temp_vault_service.get_case(case["case_id"])
    assert fetched is not None
    assert fetched["docket_number"] == "Crim. Case No. 2026-10492"

    # 4. Add Document (TXT / Markdown)
    sample_text = """
    JUDICIAL AFFIDAVIT OF WITNESS MARIA SANTOS
    I, Maria Santos, of legal age, Filipino, depose and state:
    1. Q: Do you know the accused Juan Dela Cruz?
       A: Yes, he posted the defamatory statement on Facebook on October 12, 2025.
    2. Q: What was the exact post?
       A: He accused our company of fraudulent transactions without basis.
    """
    doc_entry = temp_vault_service.add_document_to_case(
        case_id=case["case_id"],
        filename="Judicial_Affidavit_Santos.txt",
        file_bytes=sample_text.encode("utf-8"),
        doc_category="Judicial Affidavit"
    )
    assert doc_entry["filename"] == "Judicial_Affidavit_Santos.txt"
    assert doc_entry["chunk_count"] >= 1

    # Verify updated case
    updated = temp_vault_service.get_case(case["case_id"])
    assert updated["document_count"] == 1

    # 5. Delete Case
    deleted = temp_vault_service.delete_case(case["case_id"])
    assert deleted is True
    assert temp_vault_service.get_case(case["case_id"]) is None

def test_document_text_extraction():
    # Plain text extraction
    txt_bytes = "In the Matter of Petition for Habeas Corpus".encode("utf-8")
    assert "Habeas Corpus" in extract_document_text("test.txt", txt_bytes)

    # DOCX extraction with mock or real docx
    try:
        from docx import Document
        import io
        doc = Document()
        doc.add_paragraph("Section 1: General Provisions of the Contract")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Party A"
        table.rows[0].cells[1].text = "Party B"
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        extracted = extract_text_from_docx(buf.read())
        assert "General Provisions" in extracted
        assert "Party A" in extracted
    except Exception as e:
        pytest.fail(f"DOCX extraction test failed: {e}")
